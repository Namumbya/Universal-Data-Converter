import io
import json
import zipfile
from pathlib import Path
from xml.etree import ElementTree as ET

import pandas as pd
import pytest

import converter
from converter import Table, export_tables, load_tables

ROOT = Path(__file__).resolve().parents[1]
SAMPLES = ROOT / "samples"


def test_sample_csv_round_trips_to_json_and_xml():
    tables, warnings = load_tables((SAMPLES / "sample.csv").read_bytes(), ".csv")

    assert warnings == []
    assert tables[0].frame["name"].tolist() == ["Alice", "Bob", "Carol"]

    payload = json.loads(export_tables(tables, "JSON").data)
    assert payload[1]["name"] == "Bob"
    assert payload[1]["amount"] == 9.0

    xml_export = export_tables(tables, "XML")
    ET.fromstring(xml_export.data)
    reloaded, _ = load_tables(xml_export.data, ".xml")
    assert reloaded[0].frame["name"].tolist() == ["Alice", "Bob", "Carol"]
    assert reloaded[0].frame["amount"].tolist() == ["10.5", "9.0", "12.75"]


def test_semicolon_and_tab_delimiters_are_detected():
    tables, warnings = load_tables(b"id;name;amount\n1;Alice;10.5\n", "csv")
    assert tables[0].frame.loc[0, "name"] == "Alice"
    assert any(";" in warning for warning in warnings)

    tabbed, _ = load_tables(b"id\tname\n1\tAda\n", ".csv")
    assert tabbed[0].frame.loc[0, "name"] == "Ada"


def test_csv_encodings():
    bom, _ = load_tables("name\nAdèle\n".encode("utf-8-sig"), ".csv")
    assert bom[0].frame.loc[0, "name"] == "Adèle"

    latin, _ = load_tables("name,city\nAda,Café\n".encode("cp1252"), ".csv")
    assert latin[0].frame.loc[0, "city"] == "Café"


def test_sample_json_and_json_shapes():
    tables, _ = load_tables((SAMPLES / "sample.json").read_bytes(), ".json")
    assert tables[0].frame["name"].tolist() == ["Alice", "Bob"]

    columns, _ = load_tables(json.dumps({"id": [1, 2], "name": ["A", "B"]}).encode(), ".json")
    assert columns[0].frame["name"].tolist() == ["A", "B"]

    nested, _ = load_tables(json.dumps({"user": {"id": 1, "name": "Ada"}}).encode(), ".json")
    assert nested[0].frame.loc[0, "user.name"] == "Ada"

    lines, warnings = load_tables(
        b'{"id": 1, "name": "Ada"}\nnot-json\n{"id": 2, "name": "Bea"}\n',
        ".json",
    )
    assert lines[0].frame["name"].tolist() == ["Ada", "Bea"]
    assert warnings


def test_sample_xml_reads_nested_records_not_an_empty_column():
    tables, warnings = load_tables((SAMPLES / "sample.xml").read_bytes(), ".xml")

    assert warnings == []
    frame = tables[0].frame
    assert frame["name"].tolist() == ["Alice", "Bob"]
    assert frame["amount"].tolist() == ["10.5", "9.0"]
    assert list(frame.columns) == ["id", "name", "amount"]


def test_xml_attributes_become_columns():
    raw = b"""<root>
      <item id="1"><name>Ada</name></item>
      <item id="2"><name>Bea</name></item>
    </root>"""
    tables, _ = load_tables(raw, ".xml")
    frame = tables[0].frame
    assert frame["name"].tolist() == ["Ada", "Bea"]
    assert frame["@id"].tolist() == ["1", "2"]


def test_xml_export_uses_legal_tags_and_omits_missing_values():
    frame = pd.DataFrame(
        [{"First Name": "Alice", "1amount": 10, "a/b": "x", "note": float("nan")}]
    )
    exported = export_tables([Table("Data", frame)], "XML")
    ET.fromstring(exported.data)

    reloaded, _ = load_tables(exported.data, ".xml")
    row = reloaded[0].frame.iloc[0]
    assert row["First_Name"] == "Alice"
    assert row["f_1amount"] == "10"
    assert row["a_b"] == "x"
    assert row["note"] == ""


def test_xml_with_entity_declaration_is_rejected():
    payload = b"""<?xml version="1.0"?>
    <!DOCTYPE foo [<!ENTITY x "boom">]>
    <root><item><name>&x;</name></item></root>"""
    with pytest.raises(ValueError, match="XML"):
        load_tables(payload, ".xml")


def test_excel_sheet_names_survive_excel_json_and_csv_zip():
    buffer = io.BytesIO()
    with pd.ExcelWriter(buffer, engine="openpyxl") as writer:
        pd.DataFrame({"name": ["Ada"]}).to_excel(writer, sheet_name="People", index=False)
        pd.DataFrame({"sku": ["A-1"]}).to_excel(writer, sheet_name="Items", index=False)

    tables, _ = load_tables(buffer.getvalue(), ".xlsx")
    assert [table.name for table in tables] == ["People", "Items"]

    excel = export_tables(tables, "Excel (.xlsx)")
    reloaded, _ = load_tables(excel.data, ".xlsx")
    assert [table.name for table in reloaded] == ["People", "Items"]
    assert reloaded[1].frame.loc[0, "sku"] == "A-1"

    payload = json.loads(export_tables(tables, "JSON").data)
    assert payload[0]["sheet"] == "People"
    assert payload[1]["rows"][0]["sku"] == "A-1"

    zipped = export_tables(tables, "CSV")
    assert zipped.extension == "zip"
    with zipfile.ZipFile(io.BytesIO(zipped.data)) as archive:
        assert set(archive.namelist()) == {"People.csv", "Items.csv"}
        assert b"Ada" in archive.read("People.csv")


def test_missing_empty_and_unknown_files():
    with pytest.raises(ValueError, match="empty"):
        load_tables(b"", ".csv")
    with pytest.raises(ValueError, match="Unsupported"):
        load_tables(b"abc", ".txt")
    with pytest.raises(ValueError, match="JSON"):
        load_tables(b"[][]", ".json")


def test_pdf_tables_headers_and_page_cap(monkeypatch):
    class Page:
        def __init__(self, number):
            self.number = number

        def extract_tables(self):
            seen.append(self.number)
            if self.number == 1:
                return [[["id", "id", None], ["1", "2", "x"]]]
            return []

        def extract_text(self):
            return f"page {self.number}"

    class Document:
        pages = [Page(number) for number in range(1, 31)]

        def __enter__(self):
            return self

        def __exit__(self, *_args):
            return False

    seen = []
    monkeypatch.setattr(converter.pdfplumber, "open", lambda *_args, **_kwargs: Document())

    tables, warnings = load_tables(b"%PDF", ".pdf")
    assert seen == list(range(1, 26))
    assert list(tables[0].frame.columns) == ["id", "id_2", "column_3"]
    assert tables[0].frame.iloc[0].tolist() == ["1", "2", "x"]
    assert any("25" in warning for warning in warnings)


def test_pdf_without_tables_exports_text(monkeypatch):
    class Page:
        def extract_tables(self):
            return []

        def extract_text(self):
            return "Hello from the PDF"

    class Document:
        pages = [Page()]

        def __enter__(self):
            return self

        def __exit__(self, *_args):
            return False

    monkeypatch.setattr(converter.pdfplumber, "open", lambda *_args, **_kwargs: Document())
    tables, warnings = load_tables(b"%PDF", ".pdf")
    assert tables[0].frame.loc[0, "text"] == "Hello from the PDF"
    assert any("No tables" in warning for warning in warnings)


def test_markdown_round_trip_matches_markitdown_tables():
    tables, _ = load_tables((SAMPLES / "sample.csv").read_bytes(), ".csv")
    exported = export_tables(tables, "Markdown")
    text = exported.data.decode("utf-8")
    assert text.startswith("| id | name | amount |")
    assert "| --- | --- | --- |" in text

    reloaded, warnings = load_tables(exported.data, ".md")
    assert warnings == []
    assert reloaded[0].frame["name"].tolist() == ["Alice", "Bob", "Carol"]

    sample, _ = load_tables((SAMPLES / "sample.md").read_bytes(), ".markdown")
    assert sample[0].frame.loc[1, "name"] == "Bob"


def test_markdown_escapes_pipes_and_keeps_multiple_sheets():
    people = Table("People", pd.DataFrame([{"note": "a|b"}]))
    items = Table("Items", pd.DataFrame([{"sku": "A-1"}]))
    exported = export_tables([people, items], "Markdown")
    assert r"a\|b" in exported.data.decode("utf-8")

    reloaded, _ = load_tables(exported.data, ".md")
    assert [table.name for table in reloaded] == ["People", "Items"]
    assert reloaded[0].frame.loc[0, "note"] == "a|b"
    assert reloaded[1].frame.loc[0, "sku"] == "A-1"


def test_output_choices_include_markdown_and_pdf():
    assert "Markdown (.md)" in converter.SUPPORTED_OUTPUTS
    assert "PDF (.pdf)" in converter.SUPPORTED_OUTPUTS
    assert "Word (.docx)" in converter.SUPPORTED_OUTPUTS
    assert ".docx" in converter.SUPPORTED_INPUTS


def test_csv_exports_to_pdf():
    tables, _ = load_tables((SAMPLES / "sample.csv").read_bytes(), ".csv")
    exported = export_tables(tables, "PDF (.pdf)")
    assert exported.extension == "pdf"
    assert exported.data.startswith(b"%PDF")


def test_word_pdf_keeps_the_document_text():
    from docx import Document
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn

    document = Document()
    document.add_heading("1.1 Requirements Definition — Definition of Ready", level=2)
    item = document.add_paragraph("Business objective is clearly defined")
    properties = item._p.get_or_add_pPr()
    numbering = OxmlElement("w:numPr")
    level = OxmlElement("w:ilvl")
    level.set(qn("w:val"), "0")
    number = OxmlElement("w:numId")
    number.set(qn("w:val"), "1")
    numbering.append(level)
    numbering.append(number)
    properties.append(numbering)
    section = document.add_table(rows=2, cols=1)
    section.rows[0].cells[0].text = "3.1 Python"
    section.rows[1].cells[0].text = "Ruff — lint and format"
    buffer = io.BytesIO()
    document.save(buffer)
    raw = buffer.getvalue()

    preview = converter.docx_preview_text(raw)
    assert "Requirements Definition — Definition of Ready" in preview
    assert "• Business objective is clearly defined" in preview
    assert "3.1 Python" in preview
    assert "Ruff — lint and format" in preview

    exported = converter.docx_to_pdf(raw)
    import pdfplumber

    with pdfplumber.open(io.BytesIO(exported.data)) as pdf:
        text = "\n".join(page.extract_text() or "" for page in pdf.pages)
    assert "—" in text
    assert "3.1 Python" in text
    assert "Ruff — lint and format" in text
    assert "Business objective is clearly defined" in text


def test_bundled_fonts_are_in_the_repo():
    fonts = ROOT / "fonts"
    assert (fonts / "DejaVuSans.ttf").is_file()
    assert (fonts / "DejaVuSans-Bold.ttf").is_file()


def test_word_document_converts_to_pdf():
    from docx import Document

    document = Document()
    document.add_heading("Report", level=1)
    document.add_paragraph("Hello from Word")
    table = document.add_table(rows=2, cols=2)
    table.rows[0].cells[0].text = "name"
    table.rows[0].cells[1].text = "city"
    table.rows[1].cells[0].text = "Ada"
    table.rows[1].cells[1].text = "Paris"
    buffer = io.BytesIO()
    document.save(buffer)
    raw = buffer.getvalue()

    tables, _ = load_tables(raw, ".docx")
    assert tables[0].frame.loc[0, "name"] == "Ada"
    assert tables[0].frame.loc[0, "city"] == "Paris"

    exported = converter.docx_to_pdf(raw)
    assert exported.extension == "pdf"
    assert exported.data.startswith(b"%PDF")
    import pdfplumber

    with pdfplumber.open(io.BytesIO(exported.data)) as pdf:
        text = "\n".join(page.extract_text() or "" for page in pdf.pages)
    assert "Hello from Word" in text
    assert "Ada" in text


def test_pdf_and_word_convert_both_ways():
    tables, _ = load_tables((SAMPLES / "sample.csv").read_bytes(), ".csv")
    word_export = export_tables(tables, "Word (.docx)")
    assert word_export.extension == "docx"
    as_word, _ = load_tables(word_export.data, ".docx")
    assert as_word[0].frame["name"].tolist() == ["Alice", "Bob", "Carol"]

    word_pdf = converter.docx_to_pdf((SAMPLES / "sample.docx").read_bytes())
    assert word_pdf.extension == "pdf"
    import pdfplumber

    with pdfplumber.open(io.BytesIO(word_pdf.data)) as pdf:
        pdf_text = "\n".join(page.extract_text() or "" for page in pdf.pages)
    assert "Alice" in pdf_text and "Carol" in pdf_text


def test_markdown_without_a_table_is_rejected():
    with pytest.raises(ValueError, match="Markdown"):
        load_tables(b"# Notes\n\nJust a paragraph.\n", ".md")


def test_oversized_and_zip_bomb_files_are_rejected(monkeypatch):
    monkeypatch.setattr(converter, "MAX_UPLOAD_BYTES", 10)
    with pytest.raises(ValueError, match="50 MB"):
        load_tables(b"a,b\n1,2\n3,4\n", ".csv")
    monkeypatch.undo()

    monkeypatch.setattr(converter, "MAX_UNPACKED_BYTES", 100)
    buffer = io.BytesIO()
    with zipfile.ZipFile(buffer, "w", zipfile.ZIP_DEFLATED) as archive:
        archive.writestr("big.txt", "0" * 10_000)
    with pytest.raises(ValueError, match="expands"):
        load_tables(buffer.getvalue(), ".xlsx")


def test_csv_and_excel_exports_neutralize_formulas():
    frame = pd.DataFrame(
        {"v": ["=1+1", "@SUM(A1)", "+cmd", "-5", "-1e-3", "-x", "plain", "\t=1+1", "＝1"]}
    )
    values = pd.read_csv(io.BytesIO(export_tables([Table("T", frame)], "CSV").data))["v"].tolist()
    assert values == ["'=1+1", "'@SUM(A1)", "'+cmd", "-5", "-1e-3", "'-x", "plain", "'\t=1+1", "'＝1"]

    reloaded, _ = load_tables(export_tables([Table("T", frame)], "Excel (.xlsx)").data, ".xlsx")
    assert reloaded[0].frame["v"].tolist()[0] == "'=1+1"
    assert float(reloaded[0].frame["v"].tolist()[3]) == -5


def test_xml_export_strips_control_characters():
    frame = pd.DataFrame({"note": ["ab"]})
    exported = export_tables([Table("T", frame)], "XML")
    reloaded, _ = load_tables(exported.data, ".xml")
    assert reloaded[0].frame.loc[0, "note"] == "ab"


def test_corrupt_files_give_friendly_errors():
    with pytest.raises(ValueError, match="PDF"):
        load_tables(b"not a pdf", ".pdf")
    with pytest.raises(ValueError, match="Word"):
        load_tables(b"not a docx", ".docx")
    with pytest.raises(ValueError, match="Excel"):
        load_tables(b"not an xlsx", ".xlsx")


def test_word_without_tables_falls_back_to_text():
    from docx import Document

    document = Document()
    document.add_paragraph("First line")
    document.add_paragraph("Second line")
    buffer = io.BytesIO()
    document.save(buffer)

    tables, warnings = load_tables(buffer.getvalue(), ".docx")
    assert tables[0].frame["text"].tolist() == ["First line", "Second line"]
    assert any("No tables" in warning for warning in warnings)


def test_word_merged_cells_are_not_duplicated():
    from docx import Document

    document = Document()
    table = document.add_table(rows=2, cols=2)
    table.rows[0].cells[0].text = "a"
    table.rows[0].cells[1].text = "b"
    merged = table.rows[1].cells[0].merge(table.rows[1].cells[1])
    merged.text = "only once"
    buffer = io.BytesIO()
    document.save(buffer)

    tables, _ = load_tables(buffer.getvalue(), ".docx")
    assert tables[0].frame.iloc[0].tolist() == ["only once", ""]


def test_pdf_notes_pages_without_tables(monkeypatch):
    class Page:
        def __init__(self, with_table):
            self.with_table = with_table

        def extract_tables(self):
            return [[["a", "b"], ["1", "2"]]] if self.with_table else []

        def extract_text(self):
            return "text"

    class Document:
        pages = [Page(True), Page(False)]

        def __enter__(self):
            return self

        def __exit__(self, *_args):
            return False

    monkeypatch.setattr(converter.pdfplumber, "open", lambda *_a, **_k: Document())
    tables, warnings = load_tables(b"%PDF", ".pdf")
    assert list(tables[0].frame.columns) == ["a", "b"]
    assert tables[1].name == "Page2_Text"
    assert tables[1].frame.loc[0, "text"] == "text"
    assert any("exported as text: 2" in warning for warning in warnings)


def test_pdf_export_caps_rows(monkeypatch):
    monkeypatch.setattr(converter, "MAX_PDF_ROWS", 3)
    frame = pd.DataFrame({"n": range(10)})
    exported = export_tables([Table("T", frame)], "PDF (.pdf)")
    assert "first 3 rows" in exported.note


def test_sheet_names_are_truncated_and_unique():
    long_name = "x" * 40
    tables = [Table(long_name, pd.DataFrame({"a": [1]})), Table(long_name, pd.DataFrame({"a": [2]}))]
    reloaded, _ = load_tables(export_tables(tables, "Excel (.xlsx)").data, ".xlsx")
    names = [table.name for table in reloaded]
    assert len(set(names)) == 2 and all(len(name) <= 31 for name in names)


def test_empty_frames_export_in_every_format():
    tables = [Table("Empty", pd.DataFrame())]
    for output in converter.SUPPORTED_OUTPUTS:
        assert export_tables(tables, output).data


def test_old_excel_xls_round_trip():
    tables, warnings = load_tables((SAMPLES / "sample.xls").read_bytes(), ".xls")
    assert warnings == []
    assert tables[0].name == "People"
    assert tables[0].frame["name"].tolist() == ["Alice", "Bob"]


def test_pdf_without_unicode_font_replaces_unknown_characters(monkeypatch):
    from docx import Document

    monkeypatch.setattr(converter, "_pdf_font_files", lambda: (None, None))
    document = Document()
    document.add_paragraph("Size — large")
    buffer = io.BytesIO()
    document.save(buffer)
    exported = converter.docx_to_pdf(buffer.getvalue())
    import pdfplumber

    with pdfplumber.open(io.BytesIO(exported.data)) as pdf:
        text = "\n".join(page.extract_text() or "" for page in pdf.pages)
    assert "—" not in text
    assert "Size ? large" in text


def test_xml_tags_for_symbols_and_xml_prefix():
    frame = pd.DataFrame([{"!!!": "a", "xmlId": "b"}])
    exported = export_tables([Table("T", frame)], "XML")
    ET.fromstring(exported.data)
    reloaded, _ = load_tables(exported.data, ".xml")
    assert list(reloaded[0].frame.columns) == ["field", "f_xmlId"]


def test_markdown_pdf_keeps_prose_as_well_as_tables():
    raw = "\n".join(
        [
            "# Thesis",
            "",
            "Phase 1 builds a hub.",
            "",
            "- Entra SSO",
            "",
            "| Service | Owns |",
            "| --- | --- |",
            "| Search | Query records |",
            "",
            "```json",
            '{"email": "ada@example.com"}',
            "```",
            "",
        ]
    ).encode()
    preview = converter.markdown_preview_text(raw)
    assert "Thesis" in preview
    assert "Phase 1 builds a hub." in preview
    assert "• Entra SSO" in preview
    assert "Search | Query records" in preview

    exported = converter.markdown_to_pdf(raw)
    import pdfplumber

    with pdfplumber.open(io.BytesIO(exported.data)) as pdf:
        text = "\n".join(page.extract_text() or "" for page in pdf.pages)
    assert "Thesis" in text
    assert "Phase 1 builds a hub." in text
    assert "Entra SSO" in text
    assert "Query records" in text
    assert "ada@example.com" in text


def test_markdown_ignores_fenced_code_headings():
    text = ("# Real", "", "```", "# fake", "```", "", "| a |", "| --- |", "| 1 |", "")
    text = "\n".join(text).encode()
    tables, _ = load_tables(text, ".md")
    assert tables[0].name == "Real"


def test_formula_headers_are_neutralized_and_reported():
    frame = pd.DataFrame({'=HYPERLINK("http://x")': ["a"], "note": ["-x"]})
    exported = export_tables([Table("T", frame)], "CSV")
    header = exported.data.decode("utf-8-sig").splitlines()[0]
    assert header.startswith("\"'=HYPERLINK")
    assert "2 cell(s)" in exported.note

    excel = export_tables([Table("T", frame)], "Excel (.xlsx)")
    assert "2 cell(s)" in excel.note
    reloaded, _ = load_tables(excel.data, ".xlsx")
    assert str(reloaded[0].frame.columns[0]).startswith("'=")


def test_placeholder_dashes_are_left_alone():
    frame = pd.DataFrame({"v": ["-", "- item", "-5", "n/a"]})
    exported = export_tables([Table("T", frame)], "CSV")
    assert exported.note == ""
    values = pd.read_csv(io.BytesIO(exported.data), keep_default_na=False)["v"].tolist()
    assert values == ["-", "- item", "-5", "n/a"]


def test_word_file_is_parsed_once_for_text_fallback(monkeypatch):
    from docx import Document

    document = Document()
    document.add_paragraph("Only text")
    buffer = io.BytesIO()
    document.save(buffer)

    calls = []
    original = converter._docx_blocks
    monkeypatch.setattr(converter, "_docx_blocks", lambda data: calls.append(1) or original(data))
    tables, _ = load_tables(buffer.getvalue(), ".docx")
    assert tables[0].frame["text"].tolist() == ["Only text"]
    assert len(calls) == 1
