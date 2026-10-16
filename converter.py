"""Convert tabular files between CSV, Excel, JSON, XML, and Markdown.

The functions here do not import Streamlit, so they can be tested on their own.
"""

from __future__ import annotations

import csv
import html
import io
import json
import re
import zipfile
from dataclasses import dataclass
from pathlib import Path
from xml.etree import ElementTree as ET

import pandas as pd
import pdfplumber
from defusedxml.ElementTree import parse as parse_xml
from defusedxml.common import DefusedXmlException

SUPPORTED_INPUTS = [".csv", ".xlsx", ".xls", ".json", ".xml", ".pdf", ".md", ".docx"]
SUPPORTED_OUTPUTS = [
    "CSV",
    "Excel (.xlsx)",
    "JSON",
    "XML",
    "Markdown (.md)",
    "PDF (.pdf)",
    "Word (.docx)",
]

MAX_UPLOAD_BYTES = 50 * 1024 * 1024
MAX_PDF_PAGES = 25
MAX_PDF_ROWS = 2000
MAX_UNPACKED_BYTES = 200 * 1024 * 1024
_SNIFF_BYTES = 64 * 1024


@dataclass
class Table:
    name: str
    frame: pd.DataFrame


@dataclass
class Export:
    data: bytes
    mime: str
    extension: str
    note: str = ""


def load_tables(source, suffix: str) -> tuple[list[Table], list[str]]:
    """Read an uploaded file into one or more tables.

    `source` may be raw bytes or a binary file object. `suffix` is the input
    extension, with or without a leading dot.
    """
    data = _read_bytes(source)
    if not data:
        raise ValueError("The file is empty.")
    if len(data) > MAX_UPLOAD_BYTES:
        raise ValueError("The file is larger than 50 MB.")

    kind = _normalize_suffix(suffix)
    if kind == ".xlsx":
        _check_unpacked_size(data)
    warnings: list[str] = []
    if kind == ".csv":
        tables = [_csv_table(data, warnings)]
    elif kind in (".xlsx", ".xls"):
        tables = _excel_tables(data)
    elif kind == ".json":
        tables = [_json_table(data, warnings)]
    elif kind == ".xml":
        tables = _xml_tables(data, warnings)
    elif kind == ".pdf":
        tables = _pdf_tables(data, warnings)
    elif kind in {".md", ".markdown"}:
        tables = _markdown_tables(data)
    elif kind == ".docx":
        tables = _docx_tables(data, warnings)
    else:
        supported = ", ".join(SUPPORTED_INPUTS)
        raise ValueError(f"Unsupported input '{kind}'. Use one of: {supported}.")

    if not tables:
        raise ValueError("No tabular data was found in the file.")
    return tables, warnings


def export_tables(tables: list[Table], output: str) -> Export:
    """Serialize tables to the selected output format."""
    if not tables:
        raise ValueError("Nothing to export.")

    kind = _normalize_output(output)
    if kind == "csv":
        return _export_csv(tables)
    if kind == "excel":
        return _export_excel(tables)
    if kind == "json":
        return _export_json(tables)
    if kind == "xml":
        return _export_xml(tables)
    if kind == "markdown":
        return _export_markdown(tables)
    if kind == "pdf":
        return _export_pdf(tables)
    if kind == "word":
        return _export_docx(tables)
    raise ValueError(f"Unsupported output format: {output}")


def _read_bytes(source) -> bytes:
    if isinstance(source, (bytes, bytearray)):
        return bytes(source)
    data = source.read()
    if isinstance(data, str):
        return data.encode("utf-8")
    return bytes(data)


def _check_unpacked_size(data: bytes) -> None:
    """Reject zip-based files that expand far beyond the upload limit."""
    try:
        with zipfile.ZipFile(io.BytesIO(data)) as archive:
            total = sum(info.file_size for info in archive.infolist())
    except zipfile.BadZipFile:
        return  # the format reader reports a friendlier error
    if total > MAX_UNPACKED_BYTES:
        raise ValueError("The file expands to more than 200 MB and was not opened.")


def _normalize_suffix(suffix: str) -> str:
    text = (suffix or "").strip().lower()
    if not text:
        return ""
    return text if text.startswith(".") else f".{text}"


def _normalize_output(output: str) -> str:
    text = (output or "").strip().lower()
    if text == "csv" or text.startswith("csv"):
        return "csv"
    if "excel" in text or text in {"xlsx", "xls"}:
        return "excel"
    if text == "json" or text.startswith("json"):
        return "json"
    if text == "xml" or text.startswith("xml"):
        return "xml"
    if text in {"markdown", "md"} or text.startswith("markdown"):
        return "markdown"
    if text.startswith("word") or "docx" in text:
        return "word"
    if text == "pdf" or text.startswith("pdf"):
        return "pdf"
    raise ValueError(f"Unsupported output format: {output}")


def _decode_text(data: bytes) -> str:
    for encoding in ("utf-8-sig", "utf-8", "cp1252"):
        try:
            return data.decode(encoding)
        except UnicodeDecodeError:
            continue
    return data.decode("latin-1")


def _csv_table(data: bytes, warnings: list[str]) -> Table:
    text = _decode_text(data)
    if not text.strip():
        raise ValueError("The CSV file is empty.")

    kwargs: dict = {}
    sample = text[:_SNIFF_BYTES]
    try:
        dialect = csv.Sniffer().sniff(sample, delimiters=",;\t|")
    except csv.Error:
        dialect = None
    if dialect is not None:
        kwargs["sep"] = dialect.delimiter
        if dialect.quotechar:
            kwargs["quotechar"] = dialect.quotechar
        if dialect.delimiter != ",":
            warnings.append(f"Detected '{dialect.delimiter}' as the CSV delimiter.")

    try:
        frame = pd.read_csv(io.StringIO(text), **kwargs)
    except pd.errors.EmptyDataError as exc:
        raise ValueError("The CSV file has no rows.") from exc
    except pd.errors.ParserError as exc:
        raise ValueError(f"Could not parse the CSV file: {exc}") from exc
    return Table("Data", frame)


def _excel_tables(data: bytes) -> list[Table]:
    try:
        workbook = pd.ExcelFile(io.BytesIO(data))
    except Exception as exc:
        raise ValueError(f"Could not read the Excel file: {exc}") from exc
    tables = []
    for sheet in workbook.sheet_names:
        tables.append(Table(str(sheet), workbook.parse(sheet)))
    if not tables:
        raise ValueError("The Excel file has no sheets.")
    return tables


def _json_table(data: bytes, warnings: list[str]) -> Table:
    text = _decode_text(data)
    if not text.strip():
        raise ValueError("The JSON file is empty.")
    try:
        payload = json.loads(text)
    except json.JSONDecodeError:
        payload, skipped = _read_json_lines(text)
        if not payload:
            raise ValueError("Could not parse JSON. Expected a JSON value or one JSON object per line.")
        if skipped:
            warnings.append(f"Skipped {skipped} line(s) that were not valid JSON.")
        else:
            warnings.append("Read the file as JSON Lines.")
    frame = _frame_from_json(payload)
    if frame.empty:
        warnings.append("The JSON value did not contain any rows.")
    return Table("Data", frame)


def _read_json_lines(text: str) -> tuple[list, int]:
    rows = []
    skipped = 0
    for line in text.splitlines():
        if not line.strip():
            continue
        try:
            rows.append(json.loads(line))
        except json.JSONDecodeError:
            skipped += 1
    return rows, skipped


def _is_sequence(value) -> bool:
    return isinstance(value, (list, tuple)) and not isinstance(value, (str, bytes))


def _frame_from_json(payload) -> pd.DataFrame:
    if isinstance(payload, list):
        if not payload:
            return pd.DataFrame()
        if all(isinstance(item, dict) for item in payload):
            return pd.json_normalize(payload)
        return pd.DataFrame({"value": payload})
    if isinstance(payload, dict):
        values = list(payload.values())
        if values and all(_is_sequence(value) for value in values):
            lengths = {len(value) for value in values}
            if len(lengths) == 1:
                return pd.DataFrame(payload)
        return pd.json_normalize(payload)
    return pd.DataFrame({"value": [payload]})


def _xml_tables(data: bytes, warnings: list[str]) -> list[Table]:
    try:
        root = parse_xml(io.BytesIO(data)).getroot()
    except (ET.ParseError, DefusedXmlException) as exc:
        raise ValueError(f"Could not parse XML: {exc}") from exc

    exported = _tables_from_dataset(root)
    if exported:
        return exported

    groups = _record_groups(root)
    if not groups:
        frame = _flat_xml_frame(root)
        if frame.empty:
            raise ValueError("The XML file has no elements to convert.")
        warnings.append("No repeating records were found. Exported each element as a row.")
        return [Table("Data", frame)]

    largest = max(len(elements) for elements in groups.values())
    if largest == 1:
        chosen = [max(groups.items(), key=lambda item: _richness(item[1][0]))]
    else:
        chosen = [(name, elements) for name, elements in groups.items() if len(elements) >= 2]

    tables = []
    used_names: set[str] = set()
    for name, elements in chosen:
        sheet = name if name not in used_names else _unique_name(name, used_names)
        used_names.add(sheet)
        tables.append(Table(sheet, _rows_to_frame(elements)))
    return tables


def _tables_from_dataset(root: ET.Element) -> list[Table] | None:
    """Read XML written by this app, including one element per sheet."""
    if _local_name(root.tag) != "dataset":
        return None
    tables = []
    for sheet in list(root):
        if _local_name(sheet.tag) != "sheet":
            return None
        rows = [child for child in list(sheet) if _local_name(child.tag) == "row"]
        if not rows:
            continue
        tables.append(Table(sheet.attrib.get("name") or "Data", _rows_to_frame(rows)))
    return tables or None


def _local_name(tag: str) -> str:
    text = str(tag)
    if text.startswith("{"):
        text = text.split("}", 1)[-1]
    return text or "field"


def _record_groups(root: ET.Element) -> dict[str, list[ET.Element]]:
    """Group repeating elements that look like table rows.

    A row is either a parent of leaf tags (`<item><id>1</id></item>`) or a
    leaf with attributes (`<item id="1" name="Ada"/>`). Bare text tags such as
    `<id>1</id>` are not rows, so they are not chosen over the records that
    contain them.
    """
    groups: dict[str, list[ET.Element]] = {}
    for element in root.iter():
        if not _is_record(element):
            continue
        groups.setdefault(_local_name(element.tag), []).append(element)
    return groups


def _is_record(element: ET.Element) -> bool:
    children = list(element)
    if children:
        return all(len(list(child)) == 0 for child in children)
    return len(element.attrib) > 0


def _richness(element: ET.Element) -> int:
    return len(list(element)) + len(element.attrib)


def _rows_to_frame(elements: list[ET.Element]) -> pd.DataFrame:
    rows = []
    for element in elements:
        row: dict[str, str] = {}
        for child in list(element):
            key = _unique_key(_local_name(child.tag), row)
            row[key] = (child.text or "").strip()
        for attr, value in element.attrib.items():
            key = _unique_key(f"@{_local_name(attr)}", row)
            row[key] = value
        rows.append(row)
    return pd.DataFrame(rows)


def _unique_key(key: str, row: dict) -> str:
    if key not in row:
        return key
    index = 2
    while f"{key}_{index}" in row:
        index += 1
    return f"{key}_{index}"


def _flat_xml_frame(root: ET.Element) -> pd.DataFrame:
    rows = []
    for element in root.iter():
        rows.append(
            {
                "tag": _local_name(element.tag),
                "text": (element.text or "").strip(),
                **{f"@{_local_name(key)}": value for key, value in element.attrib.items()},
            }
        )
    return pd.DataFrame(rows)


def _pdf_tables(data: bytes, warnings: list[str]) -> list[Table]:
    tables: list[Table] = []
    try:
        pdf = pdfplumber.open(io.BytesIO(data))
    except Exception as exc:
        raise ValueError(f"Could not read the PDF file: {exc}") from exc
    with pdf:
        page_count = len(pdf.pages)
        pages = pdf.pages[:MAX_PDF_PAGES]
        if page_count > MAX_PDF_PAGES:
            warnings.append(f"Scanned the first {MAX_PDF_PAGES} of {page_count} pages.")
        without_tables: list[int] = []
        try:
            for page_number, page in enumerate(pages, start=1):
                found = 0
                for grid in page.extract_tables() or []:
                    frame = _grid_to_frame(grid)
                    if frame is None:
                        continue
                    found += 1
                    tables.append(Table(f"Page{page_number}_Table{found}", frame))
                if not found:
                    without_tables.append(page_number)
        except Exception as exc:
            raise ValueError(f"Could not read the PDF file: {exc}") from exc
        if tables:
            text_tables = []
            text_pages: list[int] = []
            for page_number in without_tables:
                text = (pages[page_number - 1].extract_text() or "").strip()
                if not text:
                    continue
                text_pages.append(page_number)
                text_tables.append(Table(f"Page{page_number}_Text", pd.DataFrame({"text": [text]})))
            if text_pages:
                numbers = ", ".join(str(number) for number in text_pages[:10])
                more = "..." if len(text_pages) > 10 else ""
                warnings.append(f"Pages without a table were exported as text: {numbers}{more}.")
            return tables + text_tables
        texts = [(page.extract_text() or "") for page in pages]
    warnings.append("No tables were found in the PDF. Exported page text instead.")
    return [Table("Text", pd.DataFrame({"text": texts or [""]}))]


def _grid_to_frame(grid) -> pd.DataFrame | None:
    rows = [list(row) for row in grid or [] if row and any(cell not in (None, "") for cell in row)]
    if not rows:
        return None
    header = rows[0]
    body = rows[1:]
    if _looks_like_header(header) and body:
        columns = _dedupe_headers(header)
        width = len(columns)
        records = [_pad(row, width) for row in body]
        return pd.DataFrame(records, columns=columns)
    width = max(len(row) for row in rows)
    records = [_pad(row, width) for row in rows]
    columns = [f"column_{index}" for index in range(1, width + 1)]
    return pd.DataFrame(records, columns=columns)


def _looks_like_header(row) -> bool:
    values = [cell for cell in row if cell not in (None, "")]
    return bool(values) and all(isinstance(cell, str) for cell in values)


def _dedupe_headers(header) -> list[str]:
    seen: dict[str, int] = {}
    columns = []
    for index, cell in enumerate(header, start=1):
        name = "" if cell is None else str(cell).strip()
        if not name:
            name = f"column_{index}"
        count = seen.get(name, 0) + 1
        seen[name] = count
        columns.append(name if count == 1 else f"{name}_{count}")
    return columns


def _pad(row, width: int) -> list:
    values = list(row[:width])
    if len(values) < width:
        values.extend([None] * (width - len(values)))
    return values


def _markdown_tables(data: bytes) -> list[Table]:
    """Read GitHub-style pipe tables, the same shape MarkItDown writes."""
    lines = _decode_text(data).splitlines()
    tables: list[Table] = []
    heading = "Data"
    used_names: set[str] = set()
    index = 0
    while index < len(lines):
        stripped = lines[index].strip()
        if stripped.startswith(("```", "~~~")):
            fence = stripped[:3]
            index += 1
            while index < len(lines) and not lines[index].strip().startswith(fence):
                index += 1
            index += 1
            continue
        if stripped.startswith("#"):
            heading = stripped.lstrip("#").strip() or "Data"
            index += 1
            continue
        if _is_markdown_table(lines, index):
            header = _split_markdown_row(lines[index])
            index += 2
            rows: list[list[str]] = []
            while index < len(lines) and lines[index].strip().startswith("|"):
                rows.append(_split_markdown_row(lines[index]))
                index += 1
            columns = _dedupe_headers(header)
            width = len(columns)
            frame = pd.DataFrame([_pad(row, width) for row in rows], columns=columns)
            tables.append(Table(_unique_name(heading, used_names), frame))
            heading = "Data"
            continue
        index += 1
    if not tables:
        raise ValueError(
            "No Markdown tables were found. Use a pipe table with a header and a --- separator row."
        )
    return tables


def _is_markdown_table(lines: list[str], index: int) -> bool:
    if index + 1 >= len(lines) or not lines[index].strip().startswith("|"):
        return False
    return _is_markdown_separator(lines[index + 1])


def _is_markdown_separator(line: str) -> bool:
    cells = _split_markdown_row(line) if "|" in line else []
    return bool(cells) and all(re.fullmatch(r":?-{3,}:?", cell.strip()) for cell in cells)


def _split_markdown_row(line: str) -> list[str]:
    text = line.strip()
    if text.startswith("|"):
        text = text[1:]
    if text.endswith("|") and not text.endswith("\\|"):
        text = text[:-1]
    cells: list[str] = []
    current: list[str] = []
    escaped = False
    for char in text:
        if escaped:
            current.append(char)
            escaped = False
            continue
        if char == "\\":
            escaped = True
            continue
        if char == "|":
            cells.append("".join(current).strip())
            current = []
            continue
        current.append(char)
    if escaped:
        current.append("\\")
    cells.append("".join(current).strip())
    return cells


_MARKDOWN_LIST = re.compile(r"^(?:[-*+]|\d+[.)])\s+(.*)$")


def markdown_preview_text(data: bytes) -> str:
    """Readable text of a Markdown file, in the same order the PDF uses."""
    lines = []
    for kind, payload in _markdown_document_blocks(data):
        if kind == "table":
            lines.extend(" | ".join(row) for row in payload if any(cell.strip() for cell in row))
            continue
        if kind == "rule":
            lines.append("---")
            continue
        if kind == "list":
            lines.append(_as_bullet(payload))
            continue
        if payload.strip():
            lines.append(payload.strip())
    if not lines:
        raise ValueError("The Markdown file is empty.")
    return "\n\n".join(lines)


def markdown_to_pdf(data: bytes) -> Export:
    """Turn a Markdown document into a PDF of its headings, lists, text, and tables."""
    blocks = _markdown_document_blocks(data)
    if not blocks:
        raise ValueError("The Markdown file is empty.")
    pdf, family, has_bold, unicode_font = _new_pdf()
    wrote = False
    for kind, payload in blocks:
        if kind == "table":
            rows = [[_plain_markdown(cell) for cell in row] for row in payload]
            if rows and rows[0]:
                _write_pdf_table(pdf, rows, family, unicode_font)
                pdf.ln(3)
                wrote = True
            continue
        if kind == "code":
            text = payload.rstrip()
            if not text:
                continue
            pdf.set_font(family, size=9)
            pdf.multi_cell(0, 5, _pdf_text(text, unicode_font))
            pdf.ln(1)
            wrote = True
            continue
        if kind == "rule":
            pdf.ln(1)
            y_position = pdf.get_y()
            pdf.line(pdf.l_margin, y_position, pdf.w - pdf.r_margin, y_position)
            pdf.ln(2)
            continue
        role = kind if kind in _HEADING_SIZE or kind == "list" else "body"
        if _write_pdf_paragraph(pdf, _plain_markdown(payload), role, family, has_bold, unicode_font):
            wrote = True
    if not wrote:
        raise ValueError("The Markdown file has no text or tables to convert.")
    return Export(data=_pdf_bytes(pdf), mime="application/pdf", extension="pdf")


def markdown_to_docx(data: bytes) -> Export:
    """Turn a Markdown document into a Word file of its headings, lists, text, and tables."""
    from docx import Document

    blocks = _markdown_document_blocks(data)
    if not blocks:
        raise ValueError("The Markdown file is empty.")
    document = Document()
    wrote = False
    for kind, payload in blocks:
        if kind == "table":
            rows = [[_plain_markdown(cell) for cell in row] for row in payload]
            if not rows or not rows[0]:
                continue
            width = len(rows[0])
            word_table = document.add_table(rows=len(rows), cols=width)
            for row_index, row in enumerate(rows):
                for column_index, value in enumerate(_pad(row, width)):
                    word_table.rows[row_index].cells[column_index].text = (
                        "" if value is None else str(value)
                    )
            document.add_paragraph("")
            wrote = True
            continue
        if kind == "code":
            text = payload.rstrip()
            if not text:
                continue
            document.add_paragraph(text)
            wrote = True
            continue
        if kind == "rule":
            document.add_paragraph("—" * 12)
            continue
        text = _plain_markdown(payload).strip()
        if not text:
            continue
        if kind.startswith("heading"):
            level = int(kind[-1]) if kind[-1].isdigit() else 1
            document.add_heading(text, level=min(max(level, 1), 3))
            wrote = True
            continue
        if kind == "list":
            document.add_paragraph(text, style="List Bullet")
            wrote = True
            continue
        document.add_paragraph(text)
        wrote = True
    if not wrote:
        raise ValueError("The Markdown file has no text or tables to convert.")
    buffer = io.BytesIO()
    document.save(buffer)
    return Export(
        data=buffer.getvalue(),
        mime="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        extension="docx",
    )


def _markdown_document_blocks(data: bytes) -> list[tuple[str, str | list[list[str]]]]:
    lines = _decode_text(data).splitlines()
    blocks: list[tuple[str, str | list[list[str]]]] = []
    paragraph: list[str] = []

    def flush() -> None:
        text = " ".join(part.strip() for part in paragraph).strip()
        paragraph.clear()
        if text:
            blocks.append(("paragraph", text))

    index = 0
    while index < len(lines):
        stripped = lines[index].strip()
        if not stripped:
            flush()
            index += 1
            continue
        if stripped.startswith(("```", "~~~")):
            flush()
            fence = stripped[:3]
            code: list[str] = []
            index += 1
            while index < len(lines) and not lines[index].strip().startswith(fence):
                code.append(lines[index])
                index += 1
            index += 1
            blocks.append(("code", "\n".join(code).rstrip()))
            continue
        if re.fullmatch(r"(?:-{3,}|\*{3,}|_{3,})", stripped):
            flush()
            blocks.append(("rule", ""))
            index += 1
            continue
        if stripped.startswith("#"):
            flush()
            marks = len(stripped) - len(stripped.lstrip("#"))
            text = stripped[marks:].strip()
            level = min(marks, 3) or 1
            blocks.append((f"heading{level}", text))
            index += 1
            continue
        if _is_markdown_table(lines, index):
            flush()
            rows = [_split_markdown_row(lines[index])]
            index += 2
            while index < len(lines) and lines[index].strip().startswith("|"):
                rows.append(_split_markdown_row(lines[index]))
                index += 1
            blocks.append(("table", rows))
            continue
        item = _MARKDOWN_LIST.match(stripped)
        if item:
            flush()
            blocks.append(("list", item.group(1)))
            index += 1
            continue
        paragraph.append(stripped)
        index += 1
    flush()
    return blocks


def _plain_markdown(text: str) -> str:
    """Drop emphasis markers so snake_case and code ticks stay readable in the PDF."""
    plain = html.unescape(text or "")
    plain = re.sub(r"\*\*(.+?)\*\*", r"\1", plain)
    plain = re.sub(r"(?<!\*)\*(.+?)\*(?!\*)", r"\1", plain)
    return plain.replace("`", "")


def _export_markdown(tables: list[Table]) -> Export:
    parts = []
    several = len(tables) > 1
    for index, table in enumerate(tables, start=1):
        block = _frame_to_markdown(table.frame)
        if several:
            title = (table.name or f"Sheet{index}").replace("\n", " ").strip() or f"Sheet{index}"
            block = f"# {title}\n\n{block}"
        parts.append(block)
    raw = "\n\n".join(parts).rstrip() + "\n"
    return Export(data=raw.encode("utf-8"), mime="text/markdown", extension="md")


def _frame_to_markdown(frame: pd.DataFrame) -> str:
    columns = ["" if column is None else str(column) for column in frame.columns]
    if not columns:
        return ""
    header = "| " + " | ".join(_markdown_cell(column) for column in columns) + " |"
    separator = "| " + " | ".join("---" for _ in columns) + " |"
    body = []
    for record in _records(frame):
        values = [
            _markdown_cell("" if record.get(column) is None else record.get(column))
            for column in frame.columns
        ]
        body.append("| " + " | ".join(values) + " |")
    return "\n".join([header, separator, *body])


def _markdown_cell(value) -> str:
    text = "" if value is None else str(value)
    text = text.replace("\\", "\\\\").replace("|", "\\|")
    return text.replace("\r\n", " ").replace("\n", " ").strip()


_HEADING_SIZE = {"title": 20, "heading1": 16, "heading2": 14, "heading3": 13}


def docx_preview_text(data: bytes) -> str:
    """Readable text of a Word file, in the same order the PDF uses."""
    lines = []
    for kind, payload, *rest in _docx_blocks(data):
        if kind == "paragraph":
            text = (payload or "").strip()
            if not text:
                continue
            role = rest[0] if rest else "body"
            lines.append(_as_bullet(text) if role == "list" else text)
            continue
        for row in _document_rows(payload):
            if len(row) <= 1:
                lines.append(row[0])
            else:
                lines.append(" | ".join(row))
    if not lines:
        raise ValueError("The Word file has no text or tables to convert.")
    return "\n\n".join(lines)


def docx_to_pdf(data: bytes) -> Export:
    """Turn a Word document into a PDF of its headings, lists, and tables."""
    pdf, family, has_bold, unicode_font = _new_pdf()
    wrote = False
    for kind, payload, *rest in _docx_blocks(data):
        if kind == "paragraph":
            role = rest[0] if rest else "body"
            wrote = _write_pdf_paragraph(pdf, payload, role, family, has_bold, unicode_font) or wrote
        elif _write_pdf_docx_table(pdf, payload, family, has_bold, unicode_font):
            wrote = True
    if not wrote:
        raise ValueError("The Word file has no text or tables to convert.")
    return Export(data=_pdf_bytes(pdf), mime="application/pdf", extension="pdf")


def _write_pdf_paragraph(pdf, text: str, role: str, family: str, has_bold: bool, unicode_font: bool) -> bool:
    text = (text or "").strip()
    if not text:
        pdf.ln(2)
        return False
    if role in _HEADING_SIZE:
        pdf.set_font(family, "B" if has_bold else "", _HEADING_SIZE[role])
        pdf.multi_cell(0, 8, _pdf_text(text, unicode_font))
        pdf.ln(1)
        return True
    if role == "list":
        text = _as_bullet(text)
    pdf.set_font(family, size=11)
    pdf.multi_cell(0, 6, _pdf_text(text, unicode_font))
    pdf.ln(0.5)
    return True


def _as_bullet(text: str) -> str:
    stripped = text.lstrip()
    if stripped.startswith(("•", "-", "*", "–", "—")):
        return stripped
    return f"• {stripped}"


def _write_pdf_docx_table(pdf, rows: list[list[str]], family: str, has_bold: bool, unicode_font: bool) -> bool:
    rows = _document_rows(rows)
    if not rows:
        return False
    if all(len(row) <= 1 for row in rows):
        for index, row in enumerate(rows):
            role = "heading3" if index == 0 else "body"
            _write_pdf_paragraph(pdf, row[0], role, family, has_bold, unicode_font)
        pdf.ln(1)
        return True
    _write_pdf_table(pdf, rows, family, unicode_font)
    pdf.ln(3)
    return True


def _document_rows(rows: list[list[str]]) -> list[list[str]]:
    cleaned = []
    for row in rows:
        cells = [cell.strip() for cell in row]
        if any(cells):
            cleaned.append(cells)
    width = max((len(row) for row in cleaned), default=0)
    if width > 1:
        cleaned = [row + [""] * (width - len(row)) for row in cleaned]
    return cleaned


def _docx_tables(data: bytes, warnings: list[str]) -> list[Table]:
    blocks = _docx_blocks(data)
    tables = []
    number = 0
    for kind, payload, *_rest in blocks:
        if kind != "table":
            continue
        number += 1
        if not payload:
            continue
        columns = _dedupe_headers(payload[0])
        body = [_pad(row, len(columns)) for row in payload[1:]]
        tables.append(Table(f"Table{number}", pd.DataFrame(body, columns=columns)))
    if tables:
        return tables
    lines = [block[1] for block in blocks if block[0] == "paragraph" and block[1].strip()]
    if not lines:
        raise ValueError("The Word file has no text or tables.")
    warnings.append("No tables were found in the Word file. Preview shows the document text.")
    return [Table("Text", pd.DataFrame({"text": [line.strip() for line in lines]}))]


def _docx_blocks(data: bytes) -> list[tuple]:
    from docx import Document
    from docx.oxml.table import CT_Tbl
    from docx.oxml.text.paragraph import CT_P
    from docx.table import Table as DocxTable
    from docx.text.paragraph import Paragraph

    if len(data) > MAX_UPLOAD_BYTES:
        raise ValueError("The file is larger than 50 MB.")
    _check_unpacked_size(data)
    try:
        document = Document(io.BytesIO(data))
    except Exception as exc:
        raise ValueError(f"Could not read the Word file: {exc}") from exc
    blocks: list[tuple] = []
    for child in document.element.body.iterchildren():
        if isinstance(child, CT_P):
            paragraph = Paragraph(child, document)
            blocks.append(("paragraph", paragraph.text or "", _paragraph_role(paragraph)))
        elif isinstance(child, CT_Tbl):
            blocks.append(("table", _docx_rows(DocxTable(child, document))))
    return blocks


def _paragraph_role(paragraph) -> str:
    style = ""
    if paragraph.style is not None and paragraph.style.name:
        style = paragraph.style.name
    compact = style.replace(" ", "").lower()
    if compact == "title":
        return "title"
    if compact.startswith("heading"):
        digits = "".join(character for character in compact if character.isdigit()) or "1"
        return f"heading{digits}"
    properties = paragraph._p.pPr
    if properties is not None and properties.numPr is not None:
        return "list"
    if compact.startswith("list"):
        return "list"
    return "body"


def _docx_rows(table) -> list[list[str]]:
    seen: set = set()
    rows = []
    for row in table.rows:
        values = []
        for cell in row.cells:
            # Merged cells come back once per grid position; keep the text once.
            key = cell._tc
            values.append("" if key in seen else cell.text.strip())
            seen.add(key)
        rows.append(values)
    return rows


def _export_pdf(tables: list[Table]) -> Export:
    pdf, family, _has_bold, unicode_font = _new_pdf()
    several = len(tables) > 1
    capped = False
    for index, table in enumerate(tables, start=1):
        if several:
            pdf.set_font(family, size=14)
            pdf.multi_cell(0, 8, _pdf_text(table.name or f"Sheet {index}", unicode_font))
            pdf.ln(2)
        frame = table.frame
        if len(frame) > MAX_PDF_ROWS:
            frame = frame.head(MAX_PDF_ROWS)
            capped = True
        _write_pdf_table(pdf, _frame_rows(frame), family, unicode_font)
        pdf.ln(4)
    note = f"PDF output is limited to the first {MAX_PDF_ROWS} rows of each table." if capped else ""
    return Export(data=_pdf_bytes(pdf), mime="application/pdf", extension="pdf", note=note)


def _export_docx(tables: list[Table]) -> Export:
    from docx import Document

    document = Document()
    several = len(tables) > 1
    for index, table in enumerate(tables, start=1):
        frame = table.frame
        if several:
            document.add_heading(table.name or f"Sheet {index}", level=1)
        if list(frame.columns) == ["text"]:
            for value in frame["text"].tolist():
                document.add_paragraph("" if value is None else str(value))
            continue
        rows = _frame_rows(frame)
        if not rows or not rows[0]:
            continue
        word_table = document.add_table(rows=len(rows), cols=len(rows[0]))
        for row_index, row in enumerate(rows):
            for column_index, value in enumerate(row):
                word_table.rows[row_index].cells[column_index].text = "" if value is None else str(value)
    buffer = io.BytesIO()
    document.save(buffer)
    return Export(
        data=buffer.getvalue(),
        mime="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        extension="docx",
    )


def _frame_rows(frame: pd.DataFrame) -> list[list[str]]:
    columns = ["" if column is None else str(column) for column in frame.columns]
    rows = [columns]
    for record in _records(frame):
        rows.append(
            [
                "" if record.get(column) is None else str(record.get(column))
                for column in frame.columns
            ]
        )
    return rows


def _write_pdf_table(pdf, rows: list[list[str]], family: str, unicode_font: bool) -> None:
    if not rows or not rows[0]:
        return
    pdf.set_font(family, size=10)
    with pdf.table(width=pdf.epw, repeat_headings=0, text_align="LEFT") as pdf_table:
        for row in rows:
            pdf_row = pdf_table.row()
            width = len(rows[0])
            values = list(row[:width])
            if len(values) < width:
                values.extend([""] * (width - len(values)))
            for value in values:
                pdf_row.cell(_pdf_text(value, unicode_font))


def _new_pdf():
    from fpdf import FPDF

    pdf = FPDF()
    pdf.set_auto_page_break(auto=True, margin=15)
    pdf.add_page()
    family, has_bold, unicode_font = _register_pdf_font(pdf)
    pdf.set_font(family, size=12)
    return pdf, family, has_bold, unicode_font


def _register_pdf_font(pdf) -> tuple[str, bool, bool]:
    regular, bold = _pdf_font_files()
    if regular is None:
        return "Helvetica", True, False
    pdf.add_font("AppSans", "", str(regular))
    has_bold = bold is not None
    if has_bold:
        pdf.add_font("AppSans", "B", str(bold))
    return "AppSans", has_bold, True


def _pdf_font_files() -> tuple[Path | None, Path | None]:
    bundled = Path(__file__).resolve().parent / "fonts"
    candidates = [
        (bundled / "DejaVuSans.ttf", bundled / "DejaVuSans-Bold.ttf"),
        (Path(r"C:\Windows\Fonts\arial.ttf"), Path(r"C:\Windows\Fonts\arialbd.ttf")),
        (
            Path("/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf"),
            Path("/usr/share/fonts/truetype/dejavu/DejaVuSans-Bold.ttf"),
        ),
        (
            Path("/usr/share/fonts/truetype/liberation/LiberationSans-Regular.ttf"),
            Path("/usr/share/fonts/truetype/liberation/LiberationSans-Bold.ttf"),
        ),
    ]
    for regular, bold in candidates:
        if regular.is_file():
            return regular, bold if bold.is_file() else None
    return None, None


def _pdf_text(value, unicode_font: bool) -> str:
    text = "" if value is None else str(value)
    text = text.replace("\r\n", "\n").replace("\r", "\n")
    if unicode_font:
        return text
    return text.encode("latin-1", errors="replace").decode("latin-1")


def _pdf_bytes(pdf) -> bytes:
    return bytes(pdf.output())


def _export_csv(tables: list[Table]) -> Export:
    if len(tables) == 1:
        data, quoted = _frame_to_csv_bytes(tables[0].frame)
        return Export(data=data, mime="text/csv", extension="csv", note=_formula_note(quoted))
    buffer = io.BytesIO()
    used: set[str] = set()
    quoted = 0
    with zipfile.ZipFile(buffer, "w", compression=zipfile.ZIP_DEFLATED) as archive:
        for table in tables:
            filename = _unique_name(_safe_filename(table.name), used) + ".csv"
            data, count = _frame_to_csv_bytes(table.frame)
            quoted += count
            archive.writestr(filename, data)
    note = "CSV holds one table, so each sheet was saved as its own CSV inside a zip file."
    return Export(
        data=buffer.getvalue(),
        mime="application/zip",
        extension="zip",
        note=f"{note} {_formula_note(quoted)}".strip(),
    )


def _formula_note(count: int) -> str:
    if not count:
        return ""
    return (
        f"{count} cell(s) or header(s) that spreadsheets could run as formulas "
        "were prefixed with an apostrophe."
    )


_FORMULA_TRIGGERS = set("=+@＝＋＠")
_NUMERIC = re.compile(r"[+-]?(?:\d+\.?\d*|\.\d+)(?:[eE][+-]?\d+)?$")


def _is_formula(value: str) -> bool:
    """True when a spreadsheet would treat the cell as a formula.

    OWASP lists =, +, -, @, tab, carriage return, line feed, and the full-width
    forms of those signs. A leading apostrophe stops Excel from evaluating the
    cell. Negative numbers such as -5 stay numbers. Excel can strip that
    apostrophe if a CSV is saved and opened again; the xlsx export keeps it.
    """
    text = value.lstrip(" \t\r\n\x00")
    if not text:
        return False
    first = text[0]
    if first in "-－":
        rest = text[1:2]
        # A lone dash or "- item" is plain text in Excel; "-cmd" is not.
        if not rest or rest.isspace():
            return False
        return _NUMERIC.fullmatch(text.replace("－", "-")) is None
    return first in _FORMULA_TRIGGERS


def _neutralize_formulas(frame: pd.DataFrame) -> tuple[pd.DataFrame, int]:
    """Prefix formula-like headers and text cells with a quote.

    Returns the safe frame and how many values were changed.
    """
    changed = 0

    def fix(value):
        nonlocal changed
        if isinstance(value, str) and _is_formula(value):
            changed += 1
            return "'" + value
        return value

    safe = frame.copy()
    for column in safe.columns:
        if safe[column].dtype == object or pd.api.types.is_string_dtype(safe[column].dtype):
            safe[column] = safe[column].map(fix)
    safe.columns = [fix(column) for column in safe.columns]
    return safe, changed


def _frame_to_csv_bytes(frame: pd.DataFrame) -> tuple[bytes, int]:
    buffer = io.StringIO()
    safe, changed = _neutralize_formulas(frame)
    safe.to_csv(buffer, index=False)
    return buffer.getvalue().encode("utf-8-sig"), changed


def _export_excel(tables: list[Table]) -> Export:
    buffer = io.BytesIO()
    used: set[str] = set()
    quoted = 0
    with pd.ExcelWriter(buffer, engine="openpyxl") as writer:
        for table in tables:
            sheet = _excel_sheet_name(table.name, used)
            safe, changed = _neutralize_formulas(table.frame)
            quoted += changed
            safe.to_excel(writer, index=False, sheet_name=sheet)
    return Export(
        data=buffer.getvalue(),
        mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        extension="xlsx",
        note=_formula_note(quoted),
    )


def _export_json(tables: list[Table]) -> Export:
    records = [_records(table.frame) for table in tables]
    if len(tables) == 1:
        payload = records[0]
    else:
        payload = [{"sheet": table.name, "rows": rows} for table, rows in zip(tables, records)]
    raw = json.dumps(payload, ensure_ascii=False, indent=2)
    return Export(data=raw.encode("utf-8"), mime="application/json", extension="json")


def _records(frame: pd.DataFrame):
    if frame.empty and len(frame.columns) == 0:
        return []
    return json.loads(frame.to_json(orient="records", date_format="iso", force_ascii=False))


def _export_xml(tables: list[Table]) -> Export:
    root = ET.Element("dataset")
    for index, table in enumerate(tables, start=1):
        sheet = ET.SubElement(root, "sheet", name=table.name or f"Sheet{index}")
        tags = _xml_tags(list(table.frame.columns))
        for record in _records(table.frame):
            row = ET.SubElement(sheet, "row")
            for tag, column in zip(tags, table.frame.columns):
                cell = ET.SubElement(row, tag)
                value = record.get(column)
                cell.text = "" if value is None else _XML_ILLEGAL.sub("", str(value))
    xml_bytes = ET.tostring(root, encoding="utf-8", xml_declaration=True)
    return Export(data=xml_bytes, mime="application/xml", extension="xml")


_XML_ILLEGAL = re.compile(r"[\x00-\x08\x0b\x0c\x0e-\x1f\ufffe\uffff]")


def _xml_tags(columns: list) -> list[str]:
    seen: dict[str, int] = {}
    tags = []
    for column in columns:
        base = _xml_tag(column)
        count = seen.get(base, 0) + 1
        seen[base] = count
        tags.append(base if count == 1 else f"{base}_{count}")
    return tags


def _xml_tag(name) -> str:
    text = "" if name is None else str(name).strip()
    text = re.sub(r"[^A-Za-z0-9_.-]", "_", text)
    text = text.strip("._-") or "field"
    if text.lower().startswith("xml"):
        text = f"f_{text}"
    if not (text[0].isalpha() or text[0] == "_"):
        text = f"f_{text}"
    return text


def _excel_sheet_name(name: str, used: set[str]) -> str:
    cleaned = re.sub(r"[:\\/?*\[\]]", "_", str(name)).strip() or "Sheet"
    return _unique_name(cleaned[:31], used, limit=31)


def _safe_filename(name: str) -> str:
    cleaned = re.sub(r'[<>:"/\\|?*]', "_", str(name)).strip() or "sheet"
    return cleaned[:80]


def _unique_name(name: str, used: set[str], limit: int | None = None) -> str:
    candidate = name or "sheet"
    if candidate not in used:
        used.add(candidate)
        return candidate
    index = 2
    while True:
        suffix = f"_{index}"
        base = candidate
        if limit is not None:
            base = candidate[: max(1, limit - len(suffix))]
        renamed = f"{base}{suffix}"
        if renamed not in used:
            used.add(renamed)
            return renamed
        index += 1
